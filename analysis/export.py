"""Safe, in-memory result exports and an auditable, JSON-compatible manifest."""

from __future__ import annotations

from datetime import date, datetime, timezone
from importlib.metadata import PackageNotFoundError, version
from io import BytesIO
import json
import re
from zipfile import ZIP_DEFLATED, ZipFile

import numpy as np
import pandas as pd

from . import __version__


def _safe_cell(value):
    # Keep actual numeric values numeric. Protect untrusted string labels/values
    # from spreadsheet formula interpretation, including leading whitespace.
    if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")):
        return "'" + value
    if isinstance(value, str) and value.startswith(("\t", "\r", "\n")):
        return "'" + value
    return value


def _safe_frame(frame):
    safe = frame.copy()
    for col in safe:
        if pd.api.types.is_object_dtype(safe[col]) or isinstance(safe[col].dtype, (pd.StringDtype, pd.CategoricalDtype)):
            safe[col] = safe[col].astype(object).map(_safe_cell)
    safe.columns = [_safe_cell(str(c)) for c in safe.columns]
    return safe


def csv_bytes(frame):
    """UTF-8 BOM CSV, with formula-like string fields escaped for spreadsheet use."""
    return _safe_frame(frame).to_csv(index=False).encode("utf-8-sig")


def _sheet_name(name, used):
    stem = re.sub(r"[\[\]:*?/\\]", "_", str(name)).strip("'")[:31] or "Results"
    candidate = stem
    counter = 2
    while candidate.casefold() in used:
        suffix = f"_{counter}"
        candidate = stem[:31-len(suffix)] + suffix
        counter += 1
    used.add(candidate.casefold())
    return candidate


def excel_bytes(tables):
    """Workbook containing supplied tables, with frozen headers and readable widths."""
    from openpyxl.styles import Alignment, Font, PatternFill

    buf = BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        used = set()
        for title, table in (tables or {"Results": pd.DataFrame({"Note": ["No tables available"]})}).items():
            sheet = _sheet_name(title, used)
            frame = _safe_frame(table)
            # Excel cannot represent timezone-aware timestamps. ISO strings preserve them.
            for col in frame:
                if isinstance(frame[col].dtype, pd.DatetimeTZDtype):
                    frame[col] = frame[col].map(lambda x: x.isoformat() if pd.notna(x) else None)
            frame.to_excel(writer, sheet_name=sheet, index=False)
            ws = writer.sheets[sheet]
            ws.freeze_panes = "A2"
            if len(frame.columns):
                ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.fill = PatternFill("solid", fgColor="184A68")
                cell.font = Font(color="FFFFFF", bold=True)
                cell.alignment = Alignment(wrap_text=True, vertical="center")
            ws.row_dimensions[1].height = 28
            for col in ws.columns:
                letter = col[0].column_letter
                width = max([len(str(c.value or "")) for c in list(col)[:101]] + [10])
                ws.column_dimensions[letter].width = min(55, width + 2)
    return buf.getvalue()


def _jsonable(value):
    if value is None or value is pd.NA or value is pd.NaT:
        return None
    if isinstance(value, dict):
        # Do not drop user variable names merely because they start with '_'.
        return {str(k): _jsonable(v) for k, v in value.items()
                if str(k) not in {"_fitted", "_model", "_design"}}
    if isinstance(value, pd.DataFrame):
        return [_jsonable(row) for row in value.to_dict("records")]
    if isinstance(value, pd.Series):
        return [_jsonable(x) for x in value.tolist()]
    if isinstance(value, (list, tuple, set, np.ndarray)):
        return [_jsonable(x) for x in value]
    if isinstance(value, (np.integer,)):
        return int(value)
    if isinstance(value, (float, np.floating)):
        return float(value) if np.isfinite(value) else None
    if isinstance(value, np.bool_):
        return bool(value)
    if isinstance(value, (date, datetime, pd.Timestamp)):
        return value.isoformat()
    if isinstance(value, (str, int, bool)):
        return value
    raise TypeError(f"Cannot serialize {type(value).__name__}; omit runtime objects from the summary.")


def json_bytes(value):
    return json.dumps(_jsonable(value), indent=2, ensure_ascii=False, allow_nan=False).encode("utf-8")


def interpretations(model):
    """Generate effect-focused descriptions without significance or causal claims."""
    if not model or not model.get("success"):
        return []
    logistic = model["info"].get("family") == "logistic"
    adjusted = len(model["terms"]["variable"].unique()) > 1
    adjustment = "adjusted" if adjusted else "unadjusted"
    condition = ", conditional on the other predictors in the model." if adjusted else ". This model contains one predictor."
    lines = []
    for row in model["terms"].to_dict("records"):
        estimate, low, high = (row[k] for k in ("estimate", "ci_lower", "ci_upper"))
        if not all(np.isfinite(x) for x in (estimate, low, high)):
            continue
        contrast = "For " + str(row["term"])
        if not row.get("reference"):
            contrast += ", per one-unit increase"
        if logistic:
            lines.append(f"{contrast}, the {adjustment} odds ratio for the selected event was {estimate:.3g} "
                         f"(95% CI {low:.3g} to {high:.3g}){condition}")
        else:
            lines.append(f"{contrast}, the {adjustment} difference in the mean outcome was {estimate:.3g} "
                         f"units (95% CI {low:.3g} to {high:.3g}){condition}")
    return lines


def _model_summary(model):
    if not model:
        return None
    return {k: model.get(k) for k in ("success", "error", "n_total", "n_complete", "n_dropped", "metrics", "warnings", "info")}


def analysis_summary(*, source, variable_types, cleaning_log, settings, bivariate=None, model=None, comparison=None):
    packages = {}
    for package in ("streamlit", "pandas", "numpy", "scipy", "statsmodels", "matplotlib", "openpyxl"):
        try:
            packages[package] = version(package)
        except PackageNotFoundError:
            packages[package] = "not installed"
    import platform

    result = {
        "application": "StudyFlow", "app_version": __version__, "schema_version": 1,
        "created_utc": datetime.now(timezone.utc).isoformat(),
        "source": source, "python": platform.python_version(), "packages": packages,
        "variable_types": variable_types, "settings": settings,
        "cleaning_log": cleaning_log,
        "group_comparison": None if not comparison else {"info": comparison.get("info"), "warnings": comparison.get("warnings"), "tests": comparison.get("tests")},
        "bivariate": None if not bivariate else {"info": bivariate.get("info"), "variables": bivariate.get("variables"), "models": {k: _model_summary(v) for k,v in bivariate.get("models", {}).items()}},
        "multivariable": _model_summary(model),
        "interpretations": interpretations(model),
        "limitations": [
            "Results assume independent observations; clustered, repeated-measures and survey designs are unsupported.",
            "Complete-case analysis can change the study population and can be biased; no imputation is performed.",
            "P-value screening is optional and exploratory. Retain prespecified confounders using scientific reasoning.",
            "Reported confidence intervals and p-values are nominal and do not account for predictor selection or multiple testing.",
            "Associations are not causal effects. Odds ratios are not risk ratios.",
            "Diagnostics and performance measures are calculated on the fitting sample, not on independent validation data.",
            "A local browser reload or server restart clears the working session. Save exports before leaving.",
        ],
    }
    return _jsonable(result)


def summary_markdown(summary):
    settings = summary.get("settings", {})
    source = summary.get("source", {})
    model = summary.get("multivariable")
    lines = ["# StudyFlow analysis summary", "", f"Created (UTC): {summary.get('created_utc', '')}",
             "", "## Source and reproducibility", "", f"Source file: {source.get('name', 'not specified')}",
             f"SHA-256: `{source.get('sha256', 'not recorded')}`", "",
             "Keep the original source file together with analysis_summary.json to reproduce the settings, type overrides, cleaning operations and model choices.",
             "", "## Model", ""]
    if model:
        lines.extend([f"Outcome: {settings.get('analysis', {}).get('outcome', model.get('info', {}).get('outcome', 'see JSON settings'))}",
                      f"Method: {model.get('info', {}).get('method', model.get('info', {}).get('family', 'not available'))}",
                      f"Fit succeeded: {model.get('success')}",
                      f"Included observations: {model.get('n_complete')} of {model.get('n_total')}; excluded: {model.get('n_dropped')}."])
        if model.get("error"):
            lines.extend(["", f"Estimation issue: {model['error']}"])
        for warning in model.get("warnings") or []:
            lines.extend(["", f"Model note: {warning}"])
        lines.extend(["", "### Effect interpretations", ""])
        lines.extend("- " + x for x in summary.get("interpretations", []))
    else:
        lines.append("A multivariable model has not been fitted in this export.")
    lines.extend(["", "## Statistical methods", "",
                  "Continuous descriptive summaries include mean/SD and median/quartiles. Categorical percentages use observed values as the denominator; missing percentages use the whole relevant group. See tables and analysis_summary.json for comparison methods and model-specific samples.",
                  "", "Categorical predictors are treatment-coded relative to the recorded reference category. Ordinal predictors are also treated as complete categorical variables in this MVP. Overall joint tests, not individual category p-values, are used for screening.",
                  "", "## Limits of interpretation", ""])
    lines.extend("- " + x for x in summary.get("limitations", []))
    lines.extend(["", "## Recorded settings", "", "```json", json_bytes(settings).decode(), "```", "",
                  "## Software versions", "", "```json", json_bytes(summary.get("packages", {})).decode(), "```", ""])
    return "\n".join(lines)


def bundle_bytes(tables, summary, log, figures=None):
    """Downloadable report bundle. Does not include the raw or cleaned dataset."""
    buf = BytesIO()
    with ZipFile(buf, "w", ZIP_DEFLATED) as archive:
        archive.writestr("results.xlsx", excel_bytes(tables))
        used = set()
        for title, frame in tables.items():
            stem = re.sub(r"[^A-Za-z0-9_.-]", "_", str(title)) or "table"
            name = stem
            i = 2
            while name.lower() in used:
                name = f"{stem}_{i}"
                i += 1
            used.add(name.lower())
            archive.writestr(f"tables/{name}.csv", csv_bytes(frame))
        archive.writestr("analysis_summary.json", json_bytes(summary))
        archive.writestr("analysis_summary.md", summary_markdown(summary))
        archive.writestr("analysis_log.csv", csv_bytes(pd.DataFrame(log)))
        for name, content in (figures or {}).items():
            safe_name = re.sub(r"[^A-Za-z0-9_.-]", "_", str(name))
            archive.writestr(f"figures/{safe_name}", content)
    return buf.getvalue()
