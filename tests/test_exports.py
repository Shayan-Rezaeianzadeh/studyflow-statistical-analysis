"""Exports preserve numeric results and prevent spreadsheet formula execution."""

import csv
from io import BytesIO, StringIO
import json
from zipfile import ZipFile

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

from analysis.export import analysis_summary, bundle_bytes, csv_bytes, excel_bytes, json_bytes, summary_markdown


def test_csv_protects_formula_like_text_but_keeps_numbers():
    frame = pd.DataFrame({"=header": ["=1+1", "+cmd", "-formula", "@func", "  =1", "\tcommand"],
                          "number": [-17.25, 0, 1.5, 2, 3, 4]})
    original = frame.copy(deep=True)
    rows = list(csv.reader(StringIO(csv_bytes(frame).decode("utf-8-sig"))))
    assert rows[0][0] == "'=header"
    assert all(row[0].startswith("'") for row in rows[1:])
    assert rows[1][1] == "-17.25"
    pd.testing.assert_frame_equal(frame, original)


def test_excel_formula_protection_and_numeric_cells():
    content = excel_bytes({"Results": pd.DataFrame({"label": ["=1+1", "safe"], "estimate": [-2.75, 1.25]})})
    workbook = load_workbook(BytesIO(content), data_only=False)
    sheet = workbook["Results"]
    assert sheet["A2"].data_type == "s"
    assert sheet["A2"].value == "'=1+1"
    assert sheet["B2"].data_type == "n"
    assert sheet["B2"].value == -2.75
    assert sheet.freeze_panes == "A2"


def test_excel_sheet_names_are_valid_and_case_insensitively_unique():
    names = ["bad/name*with?characters", "A" * 40, "a" * 40, "[]:*?/\\"]
    workbook = load_workbook(BytesIO(excel_bytes({name: pd.DataFrame({"value": [1]}) for name in names})))
    assert len(workbook.sheetnames) == 4
    assert len({name.casefold() for name in workbook.sheetnames}) == 4
    assert all(len(name) <= 31 and not any(ch in name for ch in "[]:*?/\\") for name in workbook.sheetnames)


def test_json_is_standard_and_omits_private_runtime_objects():
    content = json_bytes({
        "number": np.int64(4), "flag": np.bool_(True), "missing": np.nan,
        "infinite": np.inf, "pd_missing": pd.NA,
        "date": pd.Timestamp("2025-01-02"),
        "table": pd.DataFrame({"value": [1.0, np.nan]}),
        "variable_types": {"_age": "continuous"},
        "_fitted": object(),
    })
    parsed = json.loads(content, parse_constant=lambda value: pytest.fail(f"Nonstandard JSON constant: {value}"))
    assert parsed["number"] == 4
    assert parsed["flag"] is True
    assert parsed["missing"] is None and parsed["infinite"] is None and parsed["pd_missing"] is None
    assert parsed["table"][1]["value"] is None
    assert "_fitted" not in parsed
    assert parsed["variable_types"] == {"_age": "continuous"}


def test_summary_contains_reproducibility_metadata_and_no_raw_data():
    summary = analysis_summary(source={"name": "synthetic.csv", "sha256": "abc"},
                                variable_types={"x": "continuous"}, cleaning_log=[],
                                settings={"analysis": {"outcome": "x"}})
    assert summary["source"]["sha256"] == "abc"
    assert "python" in summary and "statsmodels" in summary["packages"]
    assert summary["variable_types"] == {"x": "continuous"}
    assert "synthetic.csv" in summary_markdown(summary)
    assert "data" not in summary and "raw_data" not in summary


def test_bundle_contains_safe_artifact_names_and_no_input_dataset():
    summary = analysis_summary(source={"name": "private.csv"}, variable_types={}, cleaning_log=[], settings={})
    content = bundle_bytes({"../../table": pd.DataFrame({"estimate": [2.5]})}, summary, [],
                            figures={"../../figure.svg": b"<svg></svg>"})
    with ZipFile(BytesIO(content)) as archive:
        names = archive.namelist()
        assert "results.xlsx" in names
        assert "analysis_summary.json" in names
        assert "analysis_log.csv" in names
        assert all(".." not in name.split("/") for name in names)
        assert "private.csv" not in names
        assert json.loads(archive.read("analysis_summary.json"))["source"]["name"] == "private.csv"
